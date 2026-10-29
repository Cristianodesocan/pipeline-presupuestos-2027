"""Insights automáticos (§7.4) y tablas estilo APA comunes a los cuatro scripts de análisis.

Cada insight devuelve (título, estatus, markdown). Estatus (convención del vault):
  verificado  descripción de cifras que casan con el PDF (euros corrientes)
  estimacion  usa el deflactor de un año con IPC incompleto o previsto (constantes que tocan 2026-2027)
  hipotesis   lectura sobre la intención del gobierno: nunca la genera el pipeline, la añade la persona analista
"""
import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side

FACTORES = ("COVID (2020-2021); fondos MRR/REACT-EU (2022-2026, aislables solo por sección); inflación "
            "(2022-2023); entregas a cuenta del sistema de financiación; cambios de perímetro (FDCAN en 2017, "
            "Parlamento fuera de las secciones en 2022).")


def num(v, dec=1) -> str:
    if v is None or (isinstance(v, float) and not np.isfinite(v)):
        return "—"
    s = f"{v:,.{dec}f}"
    return s.replace(",", "X").replace(".", ",").replace("X", ".")


def meur(v) -> str:
    return num(v / 1e6, 1) if v == v else "—"


def tabla_md(df: pd.DataFrame) -> str:
    if df.empty:
        return "_Sin filas._\n"
    cab = "| " + " | ".join(str(c) for c in df.columns) + " |\n|" + "---|" * len(df.columns) + "\n"
    celda = lambda x: num(x, 2) if isinstance(x, (float, np.floating)) else ("—" if x is None else str(x))  # noqa: E731
    return cab + "".join("| " + " | ".join(celda(x) for x in fila) + " |\n" for fila in df.itertuples(index=False))


def valor(t: pd.DataFrame, metrica: str, **filtros) -> pd.Series:
    m = t.metrica == metrica
    for k, v in filtros.items():
        m &= t[k] == v
    return t[m].set_index("id_estable").valor


def unidades(t: pd.DataFrame) -> pd.DataFrame:
    return t[t.id_estable != "__total__"]


def _estatus(t: pd.DataFrame, precios: str) -> str:
    return "verificado" if precios == "corrientes" else (
        "estimacion" if (t.estatus_dato == "estimacion").any() else "verificado")


def prioridades(t: pd.DataFrame, v: dict, n: int, precios: str = "constantes") -> tuple:
    """Insight 1: quién gana y quién pierde peso en la ventana."""
    u = unidades(t[(t.precios == precios)])
    delta = valor(u, "delta_peso_pp_ventana", ventana=v["id"])
    if delta.empty:
        return (f"Prioridades {v['id']}", "verificado", "_Sin datos comparables en la ventana._\n")
    pb = valor(u, "peso_total_pct", ventana="anual", anio=v["base"])
    pf = valor(u, "peso_total_pct", ventana="anual", anio=v["fin"])
    tc = valor(u, "tcac_ventana", ventana=v["id"])
    et = u.drop_duplicates("id_estable").set_index("id_estable").etiqueta
    d = pd.DataFrame({"delta": delta}).assign(etiqueta=et, pb=pb, pf=pf, tc=tc).sort_values("delta")
    fmt = lambda x: pd.DataFrame({"Partida": x.etiqueta, f"Peso {v['base']} (%)": x.pb.map(num),  # noqa: E731
                                  f"Peso {v['fin']} (%)": x.pf.map(num), "Δ peso (pp)": x.delta.map(lambda y: num(y, 2)),
                                  "TCAC (%)": x.tc.map(num)})
    texto = (f"Entre {v['base']} y {v['fin']} ({'ventana completa' if v['completa'] else 'ventana incompleta: falta el último presupuesto'}), "
             f"en euros {precios}.\n\n**Ganan peso**\n\n" + tabla_md(fmt(d.tail(n).iloc[::-1])) +
             "\n**Pierden peso**\n\n" + tabla_md(fmt(d.head(n))))
    return (f"Prioridades de la ventana {v['id']}: quién gana y quién pierde peso", _estatus(t, precios), texto)


def ruptura(t: pd.DataFrame, actual: str, anterior: str, n: int, peso_min: float = 0.5) -> tuple:
    """Insight 2: aceleraciones y frenazos frente a la tendencia de la legislatura anterior."""
    u = unidades(t[t.precios == "constantes"])
    r = valor(u, f"ruptura_tendencia_vs_{anterior}", ventana=actual)
    ultimo = int(u[u.ventana == "anual"].anio.max())
    peso = valor(u, "peso_total_pct", ventana="anual", anio=ultimo)
    r = r[peso.reindex(r.index) >= peso_min]
    et = u.drop_duplicates("id_estable").set_index("id_estable").etiqueta
    r = r.sort_values()
    fmt = lambda x: pd.DataFrame({"Partida": et.reindex(x.index), "Δ TCAC (pp)": x.map(num)})  # noqa: E731
    texto = (f"TCAC real de `{actual}` menos TCAC real de `{anterior}`, en partidas con peso ≥ {num(peso_min)} % en {ultimo}. "
             "Positivo = acelera respecto a la legislatura anterior.\n\n**Más aceleran**\n\n" + tabla_md(fmt(r[r > 0].tail(n).iloc[::-1])) +
             "\n**Más frenan**\n\n" + tabla_md(fmt(r[r < 0].head(n))))
    return (f"¿Ruptura o continuidad? {actual} frente a {anterior}", _estatus(t, "constantes"), texto)


def real_vs_nominal(t: pd.DataFrame, ventanas: list[dict]) -> tuple:
    """Insight 3: crecimiento nominal frente a real del total de la dimensión, por ventana."""
    tot = t[(t.id_estable == "__total__") & (t.metrica == "var_pct_ventana")]
    filas = []
    for v in ventanas:
        c = tot[(tot.ventana == v["id"]) & (tot.precios == "corrientes")].valor
        k = tot[(tot.ventana == v["id"]) & (tot.precios == "constantes")].valor
        if len(c) and len(k):
            filas.append({"Ventana": v["id"], "Años": f"{v['base']}→{v['fin']}", "Nominal (%)": num(c.iloc[0]),
                          "Real (%)": num(k.iloc[0]), "Se come la inflación (pp)": num(c.iloc[0] - k.iloc[0])})
    return ("Crecimiento nominal frente a real", "estimacion", tabla_md(pd.DataFrame(filas)))


def elasticas(t: pd.DataFrame, v: dict, n: int, peso_min: float = 0.5) -> tuple:
    """Insight 6: partidas que crecen más que el presupuesto (elasticidad > 1)."""
    u = unidades(t[t.precios == "constantes"])
    e = valor(u, "elasticidad_ventana", ventana=v["id"])
    peso = valor(u, "peso_total_pct", ventana="anual", anio=v["fin"])
    e = e[(peso.reindex(e.index) >= peso_min)].sort_values(ascending=False)
    et = u.drop_duplicates("id_estable").set_index("id_estable").etiqueta
    d = pd.DataFrame({"Partida": et.reindex(e.index), "Elasticidad": e.map(lambda x: num(x, 2))})
    texto = (f"{int((e > 1).sum())} de {len(e)} partidas con peso ≥ {num(peso_min)} % crecen más que el total "
             f"en `{v['id']}` ({v['base']}→{v['fin']}).\n\n" + tabla_md(d.head(n)))
    return ("Partidas que crecen más que el presupuesto", _estatus(t, "constantes"), texto)


def concentracion(t: pd.DataFrame, anios: list[int]) -> tuple:
    """Insight 7: tendencia del índice de Herfindahl."""
    h = t[(t.metrica == "hhi") & (t.precios == "corrientes")].set_index("anio").valor
    d = pd.DataFrame({"Año": [a for a in anios if a in h.index], "HHI": [num(h[a], 0) for a in anios if a in h.index]})
    return ("Concentración del gasto (HHI, 0-10.000)", "verificado", tabla_md(d))


def saltos(t: pd.DataFrame, desde: int) -> tuple:
    """Insight 8: saltos singulares (|var| > 40 % y peso ≥ 0,5 %) desde un año."""
    u = unidades(t[(t.precios == "corrientes") & (t.ventana == "anual")])
    f = u[(u.metrica == "flag_salto") & (u.valor == 1) & (u.anio >= desde)]
    vp = u[u.metrica == "var_pct"].set_index(["id_estable", "anio"]).valor
    d = pd.DataFrame({"Año": f.anio, "Partida": f.etiqueta,
                      "Var. (%)": [num(vp.get((i, a))) for i, a in zip(f.id_estable, f.anio)]}).sort_values("Año")
    return (f"Saltos singulares desde {desde} (titular o error: revisar)", "verificado", tabla_md(d))


def preelectoral(t: pd.DataFrame, n: int, anio: int = 2027, peso_min: float = 0.5) -> tuple:
    """Insight 11: ¿el presupuesto preelectoral se desvía de la media de la legislatura (2024-2026)?"""
    ultimo = int(t[t.ventana == "anual"].anio.max())
    if ultimo < anio:
        return ("Presupuesto preelectoral", "verificado",
                f"_Pendiente: el presupuesto {anio} aún no está publicado (último año con datos: {ultimo})._\n")
    u = unidades(t[(t.precios == "constantes") & (t.ventana == "anual")])
    vp = u[u.metrica == "var_pct"].pivot_table(index="id_estable", columns="anio", values="valor")
    peso = valor(u, "peso_total_pct", anio=anio)
    previos = [a for a in (anio - 3, anio - 2, anio - 1) if a in vp.columns]
    if anio not in vp.columns or not previos:
        return ("Presupuesto preelectoral", "verificado", "_Sin variaciones comparables._\n")
    d = pd.DataFrame({"v": vp[anio], "media": vp[previos].mean(axis=1)}).dropna()
    d = d[peso.reindex(d.index) >= peso_min]
    d["desvio"] = d.v - d.media
    et = u.drop_duplicates("id_estable").set_index("id_estable").etiqueta
    d = d.sort_values("desvio", ascending=False)
    fmt = pd.DataFrame({"Partida": et.reindex(d.index), f"Var. {anio} (%)": d.v.map(num),
                        f"Media {previos[0]}-{previos[-1]} (%)": d.media.map(num), "Desvío (pp)": d.desvio.map(num)})
    return ("Presupuesto preelectoral: desvío de 2027 frente a la media de la legislatura", "estimacion",
            "Euros constantes. Partidas con peso ≥ 0,5 %.\n\n" + tabla_md(fmt.head(n)) + "\n" + tabla_md(fmt.tail(n)))


def documento(titulo: str, intro: str, insights: list[tuple]) -> str:
    partes = [f"# {titulo}", "", intro, ""]
    for k, (tit, est, txt) in enumerate(insights, 1):
        partes += [f"## {k}. {tit}", "", f"**Estatus:** {est}", "", txt.rstrip(), ""]
    partes += ["---", "", f"**Factores de confusión** (no se atribuye causalidad al gobierno): {FACTORES}", ""]
    return "\n".join(partes)


def escribir_xlsx(ruta, tablas: list[tuple]) -> None:
    """Tablas estilo APA: «Tabla N» en negrita, título en cursiva, filete bajo la cabecera y nota."""
    wb = Workbook()
    wb.remove(wb.active)
    fino = Side(style="thin")
    for k, (titulo, df, nota) in enumerate(tablas, 1):
        ws = wb.create_sheet(f"Tabla {k}")
        ws["A1"], ws["A2"] = f"Tabla {k}", titulo
        ws["A1"].font, ws["A2"].font = Font(bold=True), Font(italic=True)
        for j, c in enumerate(df.columns, 1):
            cel = ws.cell(row=4, column=j, value=str(c))
            cel.font, cel.border = Font(bold=True), Border(top=fino, bottom=fino)
            cel.alignment = Alignment(horizontal="center", wrap_text=True)
        for i, fila in enumerate(df.itertuples(index=False), 5):
            for j, v in enumerate(fila, 1):
                ws.cell(row=i, column=j, value=None if (isinstance(v, float) and not np.isfinite(v)) else v)
        ult = 4 + len(df)
        for j in range(1, len(df.columns) + 1):
            ws.cell(row=ult, column=j).border = Border(bottom=fino)
        ws.cell(row=ult + 2, column=1, value=f"Nota. {nota}").font = Font(italic=True)
        for j, c in enumerate(df.columns, 1):
            ws.column_dimensions[ws.cell(row=4, column=j).column_letter].width = max(12, min(55, len(str(c)) + 4))
    wb.save(ruta)
